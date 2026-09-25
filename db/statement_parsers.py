"""
db/statement_parsers.py — parse bank statement files into normalised transactions.

Each parser returns a list of dicts:
    {"txn_date": "YYYY-MM-DD", "name": str, "description": str,
     "amount": float (signed; negative = spend), "ext_id": str | None}

No database access here — callers handle dedupe and persistence.
"""
from __future__ import annotations

import csv
import datetime
import hashlib
import re


def _require_pypdf():
    """Import pypdf, raising a friendly, actionable error if it's missing."""
    try:
        from pypdf import PdfReader
        return PdfReader
    except ImportError:
        raise ValueError(
            "Reading PDF statements needs the 'pypdf' library, which isn't "
            "installed for this Python.\n\nInstall it by running:\n"
            "    python -m pip install pypdf\n\n"
            "(CSV and Excel statements work without it.)")


# ── helpers ────────────────────────────────────────────────────────────────

def _clean(s) -> str:
    return (str(s).strip() if s is not None else "")


def _parse_amount(raw) -> float | None:
    """Parse '−£1,234.56' / '+£50.00' / '-2.40' / '238.15' → signed float."""
    if raw is None:
        return None
    s = str(raw).strip()
    if not s:
        return None
    neg = s.startswith("-") or s.startswith("−")
    pos = s.startswith("+")
    s = s.lstrip("+-−").replace("£", "").replace(",", "").strip()
    if not s:
        return None
    try:
        val = float(s)
    except ValueError:
        return None
    return -val if neg else val


def dedupe_hash(account_id: int, txn_date: str, amount: float,
                name: str, ext_id: str | None = None, occurrence: int = 0) -> str:
    """Stable identity for a transaction. Prefers the bank's own id when present.

    Without a bank id, `occurrence` distinguishes genuinely repeated same-day
    transactions (e.g. several identical round-ups): the Nth identical row in a
    statement gets occurrence N, so re-importing the same file still dedupes
    while distinct repeats are all kept.
    """
    if ext_id:
        key = f"{account_id}|ext|{ext_id}"
    else:
        key = f"{account_id}|{txn_date}|{amount:.2f}|{name.strip().lower()}|#{occurrence}"
    return hashlib.sha1(key.encode("utf-8")).hexdigest()


# ── Monzo (CSV and XLSX share one row mapper) ──────────────────────────────

_MONZO_HEADER = "Transaction ID"


def _monzo_date(value) -> str:
    """Monzo date → 'YYYY-MM-DD'. CSV gives 'DD/MM/YYYY'; xlsx gives a datetime."""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    s = _clean(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return s


def _monzo_row(row: dict):
    """Map one Monzo record (dict keyed by header) to a normalised txn."""
    amount = _parse_amount(row.get("Amount"))
    if amount is None:
        return None
    # Name falls back to Description then Type (some rows have a blank Name).
    name = (_clean(row.get("Name")) or _clean(row.get("Description"))
            or _clean(row.get("Type")) or "—")
    return {
        "txn_date":    _monzo_date(row.get("Date")),
        "name":        name,
        "description": _clean(row.get("Description")) or _clean(row.get("Notes and #tags")),
        "amount":      amount,
        "ext_id":      _clean(row.get("Transaction ID")) or None,
    }


def parse_monzo_csv(path: str) -> list:
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        return [t for t in (_monzo_row(r) for r in reader) if t]


def parse_monzo_xlsx(path: str) -> list:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # Pick the sheet whose header row contains the Monzo id column.
    target = None
    for ws in wb.worksheets:
        first = next(ws.iter_rows(values_only=True), ())
        if first and any(_clean(c) == _MONZO_HEADER for c in first):
            target = ws
            header = list(first)
            break
    if target is None:
        return []
    out = []
    for i, values in enumerate(target.iter_rows(values_only=True)):
        if i == 0:
            continue
        row = {header[j]: values[j] for j in range(min(len(header), len(values)))}
        t = _monzo_row(row)
        if t:
            out.append(t)
    return out


# ── Chase (PDF) ────────────────────────────────────────────────────────────

_CHASE_DATE = re.compile(r"^\d{2} [A-Z][a-z]{2} \d{4}$")
_CHASE_TYPES = {"Purchase", "Transfer", "Payment", "Direct Debit", "Interest"}
# amount then balance, e.g. "-£38.80 £443.69" (balance may be glued to a type word)
_CHASE_AMT_BAL = re.compile(
    r"([+\-−]£[\d,]+\.\d{2})\s+£[\d,]+\.\d{2}")


def parse_chase_pdf(path: str) -> list:
    PdfReader = _require_pypdf()
    reader = PdfReader(path)
    lines = []
    for page in reader.pages:
        text = page.extract_text() or ""
        for ln in text.split("\n"):
            ln = ln.strip()
            if ln:
                lines.append(ln)

    txns = []
    i = 0
    n = len(lines)
    while i < n:
        ln = lines[i]
        m_date = re.match(r"^(\d{2} [A-Z][a-z]{2} \d{4})\s+(.*)$", ln)
        if not m_date:
            i += 1
            continue
        date_str = m_date.group(1)
        rest = m_date.group(2).strip()
        if rest in ("Opening balance",) or rest.startswith("Opening balance") \
                or rest.startswith("Closing balance"):
            i += 1
            continue

        # Gather this record's text: the date line plus following lines until the
        # next date line — covers both the 1-line and 3-line layouts.
        block = [rest]
        j = i + 1
        while j < n and not _CHASE_DATE.match(lines[j].split("  ")[0][:11]) \
                and not re.match(r"^\d{2} [A-Z][a-z]{2} \d{4}", lines[j]):
            block.append(lines[j])
            j += 1
        blob = " ".join(block)

        m = _CHASE_AMT_BAL.search(blob)
        if m:
            amount = _parse_amount(m.group(1))
            # description = text before the amount, minus a trailing type word
            desc = blob[:m.start()].strip()
            desc = re.sub(r"\b(Purchase|Transfer|Payment|Direct Debit|Interest)\b\s*$",
                          "", desc).strip()
            for t in _CHASE_TYPES:
                if desc.endswith(t):
                    desc = desc[:-len(t)].strip()
            if amount is not None and desc:
                try:
                    dt = datetime.datetime.strptime(date_str, "%d %b %Y")
                    iso = dt.strftime("%Y-%m-%d")
                except ValueError:
                    iso = date_str
                txns.append({
                    "txn_date":    iso,
                    "name":        desc,
                    "description": desc,
                    "amount":      amount,
                    "ext_id":      None,
                })
        i = j if j > i else i + 1
    return txns


# ── Amex (PDF) — best effort ───────────────────────────────────────────────

def parse_amex_pdf(path: str) -> list:
    """Amex statements extract with scrambled column order, so this is a loose,
    best-effort parse. The upload preview lets the user verify before saving."""
    PdfReader = _require_pypdf()
    reader = PdfReader(path)
    text = "\n".join((pg.extract_text() or "") for pg in reader.pages)
    lines = [l.strip() for l in text.split("\n") if l.strip()]

    month = re.compile(r"^[A-Z][a-z]{2}$")   # Jun, Jul...
    txns = []
    i = 0
    while i < len(lines):
        # pattern seen: 'Mon', 'DD', 'Mon', 'DD', 'DETAILS', ... , amount somewhere
        if month.match(lines[i]) and i + 3 < len(lines) \
                and re.match(r"^\d{1,2}$", lines[i + 1]) and month.match(lines[i + 2]):
            day = lines[i + 1]
            mon = lines[i]
            details = lines[i + 4] if i + 4 < len(lines) else ""
            # search a short window for an amount
            amount = None
            for k in range(i + 4, min(i + 9, len(lines))):
                a = _parse_amount(lines[k])
                if a is not None and abs(a) < 100000:
                    amount = a
                    break
            if amount is not None and details:
                try:
                    dt = datetime.datetime.strptime(f"{day} {mon} 2026", "%d %b %Y")
                    iso = dt.strftime("%Y-%m-%d")
                except ValueError:
                    iso = f"{day} {mon}"
                txns.append({
                    "txn_date":    iso,
                    "name":        details,
                    "description": details,
                    "amount":      -abs(amount),   # card charges are spend
                    "ext_id":      None,
                })
            i += 4
        else:
            i += 1
    return txns


# ── dispatcher ─────────────────────────────────────────────────────────────

def detect_and_parse(path: str) -> tuple:
    """Return (source_label, [transactions]) by sniffing the file.

    Raises ValueError with a friendly message if nothing matches.
    """
    lower = path.lower()
    if lower.endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            head = f.readline()
        if _MONZO_HEADER in head:
            return "Monzo CSV", parse_monzo_csv(path)
        raise ValueError("Unrecognised CSV format (expected a Monzo export).")

    if lower.endswith(".xlsx"):
        return "Monzo XLSX", parse_monzo_xlsx(path)

    if lower.endswith(".pdf"):
        PdfReader = _require_pypdf()
        text = ""
        try:
            text = (PdfReader(path).pages[0].extract_text() or "")
        except Exception:
            pass
        low = text.lower()
        if "chase" in low or "j.p. morgan" in low:
            return "Chase PDF", parse_chase_pdf(path)
        if "american express" in low or "amex" in low:
            return "Amex PDF (best effort)", parse_amex_pdf(path)
        # default: try Chase-style parsing
        return "PDF (generic)", parse_chase_pdf(path)

    raise ValueError("Unsupported file type — upload a .csv, .xlsx or .pdf.")
